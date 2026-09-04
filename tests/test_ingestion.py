import pytest
from docx import Document as WordDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import Workbook

from app.ingestion import (
    DocumentPage,
    chunk_document,
    clean_text,
    extract_document_pages,
    validate_upload,
)


def test_clean_text_normalizes_whitespace() -> None:
    assert clean_text("  Policy\r\n\r\n  applies\t today. ") == "Policy\napplies today."


def test_upload_validation_rejects_unsafe_and_oversized_files() -> None:
    with pytest.raises(ValueError, match="path components"):
        validate_upload("../policy.md", 10, "text/markdown", 100)
    with pytest.raises(ValueError, match="exceeds"):
        validate_upload("policy.md", 101, "text/markdown", 100)


def test_empty_documents_are_rejected() -> None:
    with pytest.raises(ValueError, match="empty"):
        validate_upload("policy.md", 0, "text/markdown", 100)
    with pytest.raises(ValueError, match="no extractable text"):
        extract_document_pages("policy.md", b" \n\t ")


def test_chunks_preserve_page_and_chunk_metadata() -> None:
    splitter = RecursiveCharacterTextSplitter(chunk_size=20, chunk_overlap=0)
    chunks = chunk_document(
        "handbook.pdf",
        [DocumentPage(2, "Remote work is available."), DocumentPage(4, "Manager approval is required.")],
        splitter,
    )
    assert {chunk.metadata["page"] for chunk in chunks} == {2, 4}
    assert [chunk.metadata["chunk"] for chunk in chunks] == list(range(len(chunks)))


def test_docx_text_is_extracted(tmp_path) -> None:
    document = WordDocument()
    document.add_paragraph("Remote work policy")
    path = tmp_path / "policy.docx"
    document.save(path)
    pages = extract_document_pages(path.name, path.read_bytes())
    assert pages[0].page_number == 1
    assert pages[0].text == "Remote work policy"


def test_xlsx_values_are_extracted_with_sheet_metadata(tmp_path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Policies"
    worksheet.append(["Policy", "Days"])
    worksheet.append(["Remote work", 3])
    path = tmp_path / "policies.xlsx"
    workbook.save(path)
    pages = extract_document_pages(path.name, path.read_bytes())
    assert pages[0].page_number == 1
    assert "Remote work 3" in pages[0].text
